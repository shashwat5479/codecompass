from __future__ import annotations

from io import BytesIO
from pathlib import Path
from urllib.parse import quote
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Resume, User
from app.routers.auth import get_current_user
from app.config import settings
from app.skills import extract_skills_from_text

router = APIRouter(prefix="/api/resumes", tags=["resumes"])

# Resumes are stored in the database (Resume.file_data). Serverless disks are read-only or
# wiped between invocations, so nothing is written to the filesystem.


def extract_resume_text(content: bytes, extension: str) -> str:
    if extension == ".pdf":
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()

    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    paragraphs.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    )
    return "\n".join(paragraphs).strip()


def find_resume_skills(text: str) -> list[str]:
    return extract_skills_from_text(text)


@router.post("/upload")
async def upload_resume(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if not file.filename or Path(file.filename).suffix.lower() not in {".pdf", ".docx"}:
        raise HTTPException(status_code=400, detail="Unsupported file type")
    if file.size and file.size > settings.max_upload_size:
        raise HTTPException(status_code=400, detail="File too large")
    content = await file.read()
    if len(content) > settings.max_upload_size:
        raise HTTPException(status_code=400, detail="File too large")
    safe_name = Path(file.filename).name
    extension = Path(safe_name).suffix.lower()
    try:
        extracted_text = extract_resume_text(content, extension)
    except (BadZipFile, PackageNotFoundError, PdfReadError, ValueError) as error:
        raise HTTPException(status_code=400, detail="The resume file could not be parsed") from error
    if not extracted_text:
        raise HTTPException(status_code=400, detail="No selectable text found; scanned-image resumes are not supported")
    extracted_skills = find_resume_skills(extracted_text)
    resume = Resume(
        user_id=current_user.id,
        file_name=safe_name,
        storage_path="database",
        file_data=content,
        content_summary=extracted_text[:2000],
        extracted_skills=", ".join(extracted_skills),
    )
    db.add(resume)
    db.commit()
    db.refresh(resume)
    return {
        "message": "Resume uploaded and analyzed",
        "resume_id": resume.id,
        "file_name": resume.file_name,
        "extracted_skills": extracted_skills,
    }


@router.get("")
def list_resumes(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = db.query(Resume).filter(Resume.user_id == current_user.id).order_by(Resume.uploaded_at.desc()).all()
    return [{"id": row.id, "file_name": row.file_name, "uploaded_at": row.uploaded_at} for row in rows]


@router.get("/{resume_id}/download")
def download_resume(resume_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    resume = db.query(Resume).filter(Resume.id == resume_id, Resume.user_id == current_user.id).first()
    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found")
    if resume.file_data:
        content = bytes(resume.file_data)
    else:
        # Resumes uploaded before the move to database storage still live on disk.
        legacy_path = Path(resume.storage_path)
        if not legacy_path.is_file():
            raise HTTPException(status_code=404, detail="Resume file is missing")
        content = legacy_path.read_bytes()
    return Response(
        content=content,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(resume.file_name)}"},
    )


@router.get("/analyze")
def analyze_resume(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    resumes = db.query(Resume).filter(Resume.user_id == current_user.id).order_by(Resume.uploaded_at.desc()).all()
    if not resumes:
        raise HTTPException(status_code=404, detail="No resume uploaded")
    all_skills = []
    summaries = []
    for r in resumes:
        if r.content_summary:
            summaries.append(f"[{r.file_name}]: {r.content_summary}")
        if r.extracted_skills:
            for s in r.extracted_skills.split(","):
                val = s.strip()
                if val and val not in all_skills:
                    all_skills.append(val)
    return {
        "resume_id": resumes[0].id,
        "total_resumes": len(resumes),
        "summary": "\n\n".join(summaries)[:2000],
        "skills": all_skills,
    }


@router.delete("/{resume_id}")
@router.delete("/{resume_id}/")
@router.post("/{resume_id}/delete")
@router.post("/{resume_id}")
@router.post("/{resume_id}/")
def delete_resume(resume_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    from app.models import Application

    resume = db.query(Resume).filter(Resume.id == resume_id, Resume.user_id == current_user.id).first()
    if not resume:
        return {"message": "Resume already removed", "deleted_id": resume_id}

    # Detach the resume from applications first so the foreign key doesn't block deletion.
    db.query(Application).filter(Application.resume_id == resume.id).update(
        {Application.resume_id: None},
        synchronize_session=False,
    )
    legacy_path = Path(resume.storage_path or "")
    if resume.storage_path and resume.storage_path != "database":
        try:
            if legacy_path.is_file():
                legacy_path.unlink()
        except OSError:
            pass

    db.delete(resume)
    db.commit()
    return {"message": "Resume deleted successfully", "deleted_id": resume_id}
